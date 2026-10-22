"""Generate the Northbridge 15-minute demo deck and presenter guide."""

from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches as DocInches, Pt as DocPt, RGBColor as DocRGBColor
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent
SOURCE_DECK = ROOT / "Northbridge-WorkIQ-Management-Briefing-Context-Updated.pptx"
PPTX_OUT = ROOT / "Northbridge_WorkIQ_15_Minute_Demo.pptx"
DOCX_OUT = ROOT / "Northbridge_WorkIQ_15_Minute_Presenter_Guide.docx"

NAVY = RGBColor(29, 45, 74)
BLUE = RGBColor(39, 104, 174)
CYAN = RGBColor(0, 151, 167)
TEAL = RGBColor(0, 120, 111)
AMBER = RGBColor(234, 150, 23)
RED = RGBColor(196, 57, 53)
INK = RGBColor(34, 40, 49)
MUTED = RGBColor(92, 103, 116)
PALE = RGBColor(241, 245, 249)
WHITE = RGBColor(255, 255, 255)


def clear_slides(prs: Presentation) -> None:
    """Keep the source theme and masters while removing its original slides."""
    for slide_id in list(prs.slides._sldIdLst):
        relationship_id = slide_id.rId
        prs.part.drop_rel(relationship_id)
        prs.slides._sldIdLst.remove(slide_id)


def add_textbox(slide, left, top, width, height, text, size=20, color=INK,
                bold=False, font="Aptos", align=PP_ALIGN.LEFT,
                margin=0.05, valign=MSO_ANCHOR.TOP):
    shape = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(margin)
    frame.margin_right = Inches(margin)
    frame.margin_top = Inches(margin)
    frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return shape


def add_rect(slide, left, top, width, height, fill, radius=False, line=None):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, Inches(left), Inches(top), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line or fill
    if radius:
        shape.adjustments[0] = 0.08
    return shape


def add_header(slide, section, title, subtitle=None, timebox=None):
    add_textbox(slide, 0.55, 0.30, 3.6, 0.3, section.upper(), 11, CYAN, True)
    add_textbox(slide, 0.55, 0.68, 11.7, 0.65, title, 27, NAVY, True)
    if subtitle:
        add_textbox(slide, 0.57, 1.36, 11.6, 0.45, subtitle, 13, MUTED)
    if timebox:
        add_rect(slide, 11.85, 0.28, 0.88, 0.38, NAVY, True)
        add_textbox(slide, 11.85, 0.30, 0.88, 0.28, timebox, 11, WHITE, True, align=PP_ALIGN.CENTER)


def add_footer(slide, number):
    add_rect(slide, 0.55, 7.13, 12.2, 0.02, RGBColor(211, 218, 226))
    add_textbox(slide, 0.57, 7.18, 5.0, 0.2, "NORTHBRIDGE HEALTH NETWORK  |  WORK IQ", 8, MUTED, True)
    add_textbox(slide, 12.2, 7.18, 0.5, 0.2, f"{number:02d}", 8, MUTED, True, align=PP_ALIGN.RIGHT)


def add_bullets(slide, left, top, width, height, bullets, size=17, color=INK):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.05)
    for index, text in enumerate(bullets):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = text
        paragraph.level = 0
        paragraph.font.name = "Aptos"
        paragraph.font.size = Pt(size)
        paragraph.font.color.rgb = color
        paragraph.space_after = Pt(10)
        paragraph.text = f"•  {text}"
    return box


def add_card(slide, left, top, width, height, accent, title, body, metric=None):
    add_rect(slide, left, top, width, height, WHITE, True, RGBColor(215, 223, 232))
    add_rect(slide, left, top, 0.08, height, accent)
    add_textbox(slide, left + 0.25, top + 0.20, width - 0.45, 0.38, title, 15, NAVY, True)
    if metric:
        add_textbox(slide, left + 0.25, top + 0.72, width - 0.45, 0.90, metric, 24, accent, True)
        add_textbox(slide, left + 0.25, top + 1.72, width - 0.45, height - 1.92, body, 12, MUTED)
    else:
        add_textbox(slide, left + 0.25, top + 0.78, width - 0.45, height - 1.0, body, 13, MUTED)


def add_flow_node(slide, left, top, width, title, detail, color):
    add_rect(slide, left, top, width, 1.16, WHITE, True, RGBColor(210, 220, 230))
    add_rect(slide, left, top, width, 0.12, color)
    add_textbox(slide, left + 0.16, top + 0.25, width - 0.32, 0.30, title, 14, NAVY, True, align=PP_ALIGN.CENTER)
    add_textbox(slide, left + 0.16, top + 0.62, width - 0.32, 0.34, detail, 10, MUTED, align=PP_ALIGN.CENTER)


def build_presentation() -> None:
    prs = Presentation(str(SOURCE_DECK))
    clear_slides(prs)
    blank = prs.slide_layouts[6]

    # 1. Introduction
    slide = prs.slides.add_slide(blank)
    add_rect(slide, 0, 0, 13.333, 7.5, NAVY)
    add_rect(slide, 0, 0, 0.18, 7.5, CYAN)
    add_textbox(slide, 0.75, 0.65, 5.8, 0.35, "MICROSOFT WORK IQ  |  15-MINUTE DEMO", 12, RGBColor(125, 220, 229), True)
    add_textbox(slide, 0.75, 1.45, 10.9, 1.45, "From committee decisions\nto governed CAPA action", 34, WHITE, True)
    add_textbox(slide, 0.78, 3.20, 9.8, 0.75, "Northbridge Health Network\nJoint Commission readiness across meetings, email, files, Teams, and the CAPA tracker", 17, RGBColor(221, 231, 240))
    add_rect(slide, 0.78, 5.42, 3.45, 0.78, CYAN, True)
    add_textbox(slide, 0.95, 5.61, 3.10, 0.30, "CITED  •  PERMISSION-TRIMMED  •  ACTIONABLE", 10, WHITE, True, align=PP_ALIGN.CENTER)
    add_textbox(slide, 11.75, 6.85, 0.8, 0.3, "0:00–1:00", 10, WHITE, True, align=PP_ALIGN.RIGHT)

    # 2. Business problem
    slide = prs.slides.add_slide(blank)
    add_header(slide, "Business problem / use case", "One readiness answer is scattered across five systems", "Maria has the permissions, but not the hours to reconcile the evidence.", "1:00–3:00")
    sources = [
        ("MEETINGS", "Decisions + owners", BLUE),
        ("EMAIL", "Why actions stalled", AMBER),
        ("TEAMS", "Latest field updates", CYAN),
        ("FILES", "Policy + evidence", TEAL),
        ("CAPA", "Status + due dates", RED),
    ]
    for index, (title, detail, color) in enumerate(sources):
        left = 0.6 + index * 2.48
        add_flow_node(slide, left, 2.15, 2.05, title, detail, color)
        if index < len(sources) - 1:
            add_textbox(slide, left + 2.08, 2.47, 0.40, 0.30, "+", 19, MUTED, True, align=PP_ALIGN.CENTER)
    add_rect(slide, 1.10, 4.10, 11.1, 1.55, PALE, True)
    add_textbox(slide, 1.35, 4.34, 10.6, 0.42, "The question Northbridge cannot answer quickly", 13, CYAN, True, align=PP_ALIGN.CENTER)
    add_textbox(slide, 1.48, 4.86, 10.35, 0.62, "“What was decided, who owns the corrective action, what is still open, and what should be updated?”", 22, NAVY, True, align=PP_ALIGN.CENTER)
    add_footer(slide, 2)

    # 3. Existing challenges
    slide = prs.slides.add_slide(blank)
    add_header(slide, "Existing challenges", "The failure is not missing data. It is missing reconciliation.", None, "3:00–5:00")
    add_card(slide, 0.65, 1.75, 3.80, 3.70, BLUE, "FRAGMENTED EVIDENCE", "Meeting decisions, follow-up email, shared policy files, and tracker rows all tell only part of the story.", "5 signal types")
    add_card(slide, 4.77, 1.75, 3.80, 3.70, AMBER, "STALE EXECUTION", "Owners and due dates drift after meetings. Leadership sees a tracker that no longer matches the latest commitments.", "Hours of manual work")
    add_card(slide, 8.89, 1.75, 3.80, 3.70, RED, "GOVERNANCE RISK", "HR-sensitive credentialing details must remain invisible to personas without access, even inside a useful synthesis.", "Zero access widening")
    add_textbox(slide, 0.72, 5.90, 12.0, 0.55, "Business consequence: readiness decisions are slower, evidence trails are weaker, and overdue CAPAs surface too late.", 17, NAVY, True, align=PP_ALIGN.CENTER)
    add_footer(slide, 3)

    # 4. Solution overview
    slide = prs.slides.add_slide(blank)
    add_header(slide, "Solution overview", "One work graph, four capability domains, one governed loop", "The simulator preserves the Work IQ tool contract so the endpoint can later be swapped for the live tenant.", "5:00–7:00")
    nodes = [
        ("CONTEXT", "Ground across work signals", BLUE),
        ("CHAT", "Decompose + synthesize", CYAN),
        ("TOOLS", "Create / update records", TEAL),
        ("WORKSPACES", "Persist multi-step state", AMBER),
    ]
    for index, (title, detail, color) in enumerate(nodes):
        left = 0.75 + index * 3.12
        add_flow_node(slide, left, 2.00, 2.55, title, detail, color)
        if index < 3:
            add_textbox(slide, left + 2.58, 2.42, 0.50, 0.25, "→", 19, MUTED, True, align=PP_ALIGN.CENTER)
    add_rect(slide, 1.25, 3.75, 10.85, 1.70, NAVY, True)
    add_textbox(slide, 1.55, 4.00, 10.2, 0.35, "ORCHESTRATED AGENT", 13, RGBColor(125, 220, 229), True, align=PP_ALIGN.CENTER)
    add_textbox(slide, 1.55, 4.48, 10.2, 0.46, "Reconciliation agent  +  citation builder  +  human approval gate", 21, WHITE, True, align=PP_ALIGN.CENTER)
    add_textbox(slide, 1.55, 5.05, 10.2, 0.25, "MCP for structured fetch/write  |  A2A for grounded natural-language delegation", 11, RGBColor(221, 231, 240), align=PP_ALIGN.CENTER)
    add_footer(slide, 4)

    # 5. Demo flow
    slide = prs.slides.add_slide(blank)
    add_header(slide, "Agent demo", "Live demo: evidence → recommendation → approved write", "Switch to the browser at http://127.0.0.1:8000", "7:00–10:30")
    steps = [
        ("1", "ASK", "Reconcile latest meeting actions against the current CAPA tracker.", BLUE),
        ("2", "INSPECT", "Show proposed create/update operations with source citations.", CYAN),
        ("3", "APPROVE", "Human confirms the recommendation before any write occurs.", AMBER),
        ("4", "VERIFY", "Re-fetch tracker and show the new status plus audit trail.", TEAL),
    ]
    for index, (number, title, body, color) in enumerate(steps):
        top = 1.82 + index * 1.18
        add_rect(slide, 0.88, top, 0.72, 0.72, color, True)
        add_textbox(slide, 0.88, top + 0.10, 0.72, 0.40, number, 21, WHITE, True, align=PP_ALIGN.CENTER)
        add_textbox(slide, 1.90, top + 0.02, 1.35, 0.35, title, 14, color, True)
        add_textbox(slide, 3.15, top, 8.85, 0.62, body, 16, INK)
    add_rect(slide, 9.83, 6.20, 2.28, 0.46, PALE, True)
    add_textbox(slide, 9.90, 6.30, 2.14, 0.22, "Never auto-commit", 11, RED, True, align=PP_ALIGN.CENTER)
    add_footer(slide, 5)

    # 6. Governance
    slide = prs.slides.add_slide(blank)
    add_header(slide, "Agent demo", "Same question. Different identity. Correctly different answer.", "Show permission trimming live without changing code or data.", "10:30–12:00")
    add_card(slide, 0.75, 1.82, 5.70, 3.65, CYAN, "QUALITY PROGRAM MANAGER", "Useful synthesis returns, but FILE-003 and its HR-sensitive closed-session detail are withheld. A governance note explains the omission.", "quality_pm")
    add_card(slide, 6.88, 1.82, 5.70, 3.65, TEAL, "OPERATIONS DIRECTOR", "The same prompt returns the additional HR-sensitive citation and detail because this persona has explicit access.", "ops_director")
    add_textbox(slide, 0.95, 5.82, 11.5, 0.60, "Control demonstrated: source-level permission trimming under delegated identity. Work IQ never widens access.", 18, NAVY, True, align=PP_ALIGN.CENTER)
    add_footer(slide, 6)

    # 7. Technical deep dive
    slide = prs.slides.add_slide(blank)
    add_header(slide, "Technical deep dive", "How the agent stays grounded, auditable, and replaceable", None, "12:00–14:00")
    rows = [
        ("GROUNDING", "Synthetic meetings, email, Teams, files, people, CAPA JSON", BLUE),
        ("REASONING", "Foundry model orchestrates retrieval and synthesis", CYAN),
        ("TRANSPORT", "MCP structured tools + A2A grounded delegation", TEAL),
        ("CONTROL", "Persona ACL trimming + HITL proposal approval", AMBER),
        ("TRACE", "Citation IDs + created_by / modified_by audit metadata", RED),
    ]
    for index, (label, detail, color) in enumerate(rows):
        top = 1.67 + index * 0.92
        add_rect(slide, 0.80, top, 2.10, 0.60, color, True)
        add_textbox(slide, 0.88, top + 0.12, 1.94, 0.28, label, 12, WHITE, True, align=PP_ALIGN.CENTER)
        add_textbox(slide, 3.20, top + 0.05, 9.0, 0.45, detail, 16, INK)
    add_textbox(slide, 0.84, 6.44, 11.8, 0.36, "Production transition: replace the simulator endpoint; preserve the agent’s Work IQ contract and governance behavior.", 14, NAVY, True, align=PP_ALIGN.CENTER)
    add_footer(slide, 7)

    # 8. Impact and roadmap
    slide = prs.slides.add_slide(blank)
    add_header(slide, "Impact & future roadmap", "A five-minute readiness brief, with the tracker already aligned", None, "14:00–15:00")
    add_card(slide, 0.68, 1.72, 3.82, 3.60, TEAL, "TIME", "Readiness preparation collapses from hours of app-hopping to one cited conversation.", "2–3 hours → 5 min")
    add_card(slide, 4.76, 1.72, 3.82, 3.60, BLUE, "CONTROL", "Every recommendation retains source evidence, identity scope, and a human decision point.", "100% cited")
    add_card(slide, 8.84, 1.72, 3.82, 3.60, AMBER, "NEXT", "Pilot CAPA reconciliation first, then extend to vendor status and procurement evaluation.", "90-day pilot")
    add_rect(slide, 1.05, 5.80, 11.25, 0.70, NAVY, True)
    add_textbox(slide, 1.25, 5.98, 10.85, 0.30, "Decision requested: sponsor a Northbridge pilot with named users, baseline metrics, and tenant authorization.", 15, WHITE, True, align=PP_ALIGN.CENTER)
    add_footer(slide, 8)

    prs.save(PPTX_OUT)


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def set_cell_text(cell, text, bold=False, color="202832", size=9) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Aptos"
    run.font.size = DocPt(size)
    run.font.color.rgb = DocRGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_doc_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        shade_cell(table.rows[0].cells[index], "1D2D4A")
        set_cell_text(table.rows[0].cells[index], header, True, "FFFFFF", 9)
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            set_cell_text(cells[index], str(value), False, "222831", 9)
        if len(table.rows) % 2 == 1:
            for cell in cells:
                shade_cell(cell, "F1F5F9")
    if widths:
        for row in table.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = DocInches(width)
    return table


def add_doc_heading(doc, text, level=1):
    heading = doc.add_heading(text, level=level)
    heading.style.font.name = "Aptos Display"
    heading.style.font.color.rgb = DocRGBColor(29, 45, 74)
    return heading


def add_prompt(doc, label, prompt):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = DocInches(0.25)
    paragraph.paragraph_format.right_indent = DocInches(0.25)
    paragraph.paragraph_format.space_before = DocPt(3)
    paragraph.paragraph_format.space_after = DocPt(8)
    run = paragraph.add_run(f"{label}: ")
    run.bold = True
    run.font.color.rgb = DocRGBColor(0, 120, 111)
    prompt_run = paragraph.add_run(prompt)
    prompt_run.italic = True


def add_bullet(doc, text, level=0):
    paragraph = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    paragraph.add_run(text)
    return paragraph


def build_presenter_guide() -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = DocInches(0.55)
    section.bottom_margin = DocInches(0.55)
    section.left_margin = DocInches(0.65)
    section.right_margin = DocInches(0.65)
    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = DocPt(10)

    title = doc.add_heading("Northbridge Work IQ — 15-Minute Demo Presenter Guide", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = doc.add_paragraph("CAPA reconciliation, governed action, and permission trimming")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.runs[0].italic = True
    subtitle.runs[0].font.color.rgb = DocRGBColor(0, 120, 111)

    doc.add_paragraph(
        "Purpose: deliver the manager-requested 15-minute recording format while covering the hackathon’s strongest scoring areas: explicit work-context architecture, multi-signal synthesis, a governed Tools write, citations, and live permission trimming."
    )

    add_doc_heading(doc, "Run of Show")
    add_doc_table(
        doc,
        ["Time", "Slide / screen", "Purpose", "Hard stop"],
        [
            ("0:00–1:00", "1 — Introduction", "Industry, team focus, and one-sentence promise", "1:00"),
            ("1:00–3:00", "2 — Business problem", "User, use case, and five work signals", "3:00"),
            ("3:00–5:00", "3 — Existing challenges", "Fragmentation, stale execution, governance risk", "5:00"),
            ("5:00–7:00", "4 — Solution overview", "Context, Chat, Tools, Workspace; MCP + A2A", "7:00"),
            ("7:00–10:30", "5 + live browser", "CAPA reconcile, approve, write, verify", "10:30"),
            ("10:30–12:00", "6 + live browser", "Same prompt as two personas", "12:00"),
            ("12:00–14:00", "7 — Technical deep dive", "Grounding, orchestration, integrations, controls", "14:00"),
            ("14:00–15:00", "8 — Impact / roadmap", "Outcome, pilot, decision request", "15:00"),
        ],
        [1.1, 1.6, 4.0, 0.8],
    )

    add_doc_heading(doc, "Before You Record")
    add_bullet(doc, "Use the local Work IQ Simulator and say so explicitly. Do not imply this run uses a live Microsoft 365 tenant.")
    add_bullet(doc, "Restore the C1 Northbridge fixtures to a known baseline. The current capa_tracker.json contains a demo-created row with id NEW, so remove or archive that row before rehearsing the create flow.")
    add_bullet(doc, "Confirm CAPA-001 through CAPA-005 are present and record a screenshot of the pre-demo tracker state.")
    add_bullet(doc, "Start the agent UI from the repository root: .\\.venv\\Scripts\\python.exe agent\\web.py")
    add_bullet(doc, "Open http://127.0.0.1:8000 and create one quality_pm session and one ops_director session before recording.")
    add_bullet(doc, "Keep the PowerPoint, browser, capa_tracker.json, and this presenter guide open. Disable desktop notifications.")
    add_bullet(doc, "Run a full rehearsal and save screenshots of the expected proposal, permission note, and final tracker as fallback evidence.")

    add_doc_heading(doc, "Presenter Script")
    sections = [
        ("0:00–1:00 — Introduction", [
            "Northbridge Health Network runs readiness work across committees, email, Teams, shared files, and a CAPA tracker. Our Work IQ agent turns those fragmented signals into one cited answer and, with human approval, one governed tracker update.",
            "Today I will show the full loop: find the evidence, recommend an action, approve the write, and prove permissions are enforced.",
        ]),
        ("1:00–3:00 — Business Problem / Use Case", [
            "The primary user is Maria Delgado, a quality program manager preparing for a Joint Commission readiness review.",
            "Her question is simple, but the answer is not: what was decided, who owns it, why is it stalled, and does the CAPA tracker reflect the latest commitment?",
            "The answer spans meeting recaps, email rationale, Teams updates, shared evidence, and structured CAPA records. No single application can see it whole.",
        ]),
        ("3:00–5:00 — Existing Challenges", [
            "First, evidence is fragmented. Second, execution records go stale after meetings. Third, sensitive credentialing material must remain permission-trimmed.",
            "The operational cost is hours of manual reconciliation and a weak evidence trail precisely when leadership needs confidence.",
        ]),
        ("5:00–7:00 — Solution Overview", [
            "Context retrieves permission-scoped work signals. Chat decomposes the compound question. Tools fetch and update the CAPA table. Workspace represents tenant-bounded state for longer-running work.",
            "The agent uses A2A for grounded delegation and MCP for structured fetch and write operations. Citations are preserved end to end.",
            "The reconciliation agent never writes automatically: it creates a proposal, and the human approval gate decides whether the update executes.",
        ]),
    ]
    for heading, paragraphs in sections:
        add_doc_heading(doc, heading, 2)
        for text in paragraphs:
            doc.add_paragraph(text)

    add_doc_heading(doc, "7:00–10:30 — Live Demo: CAPA Reconciliation", 2)
    doc.add_paragraph("Switch to the browser. Keep narration short and point to citations, proposal status, and the tracker result.")
    add_prompt(doc, "Prompt 1", "Are there any new action items from the latest meetings that should be reflected in the CAPA tracker? Compare against existing open items and recommend creates or updates. Do not write anything yet.")
    add_bullet(doc, "Say: The agent is comparing unstructured meeting actions with existing structured records.")
    add_bullet(doc, "Point out: proposed operation, matched/new classification, owner, due date, and source citation.")
    add_bullet(doc, "Point out: proposal remains pending; no tracker write has occurred.")
    add_prompt(doc, "Prompt 2", "Approve the PharmaLink compliance-audit proposal and apply the recommended CAPA-001 update.")
    add_bullet(doc, "Say: This is the governed Tools moment. Human approval changes the proposal from recommendation to action.")
    add_prompt(doc, "Prompt 3", "Fetch the CAPA tracker and show the created or modified items with audit metadata and citations.")
    add_bullet(doc, "Point out: the write is observable on re-query, including created_by or modified_by metadata.")
    add_bullet(doc, "Do not dwell on IDs or dates if a rehearsal generates different idempotent values; demonstrate the state transition and traceability.")

    add_doc_heading(doc, "10:30–12:00 — Live Demo: Permission Trimming", 2)
    governance_prompt = "Prep me for the Joint Commission readiness review: what did the quality committee decide on medication reconciliation, who owns the corrective action, did credentialing close the onboarding gap, and what is outstanding?"
    add_prompt(doc, "As quality_pm", governance_prompt)
    add_bullet(doc, "Point out the useful multi-signal answer and the governance note withholding FILE-003.")
    add_prompt(doc, "As ops_director", governance_prompt)
    add_bullet(doc, "Point out that FILE-003 and the closed-session detail now appear because the identity has access.")
    add_bullet(doc, "Say: Same question, same corpus, no code change. Identity controls the answer at source level.")

    add_doc_heading(doc, "12:00–14:00 — Technical Deep Dive", 2)
    add_bullet(doc, "Data sources: synthetic email, meeting, Teams, file, people, and CAPA fixtures for the C1 Northbridge scenario.")
    add_bullet(doc, "Orchestration: a Foundry-backed agent chooses retrieval and tool calls; reconciliation and citation subagents specialize the workflow.")
    add_bullet(doc, "Integrations: MCP exposes ask/fetch/create/update tools; A2A exposes grounded question delegation through an agent card.")
    add_bullet(doc, "Controls: delegated persona identity, ACL trimming, explicit governance notes, approval-gated writes, citations, and audit metadata.")
    add_bullet(doc, "Transition: the simulator mirrors the contract, not the scale. Production requires Work IQ enablement, admin consent for WorkIQAgent.Ask, Copilot licensing, and user write access to the Dataverse CAPA table.")

    add_doc_heading(doc, "14:00–15:00 — Impact and Roadmap", 2)
    doc.add_paragraph("Close with the outcome, not another feature list:")
    add_bullet(doc, "Before: 2–3 hours to prepare a cross-committee readiness brief and manually reconcile the tracker.")
    add_bullet(doc, "After: approximately five minutes to obtain a cited brief and approve an auditable tracker update.")
    add_bullet(doc, "Pilot: 90 days, CAPA reconciliation first; measure preparation time, stale-record rate, citation coverage, and approval turnaround.")
    add_bullet(doc, "Decision request: sponsor named pilot users and the tenant authorization needed to connect the same contract to live Work IQ.")

    add_doc_heading(doc, "Demo Recovery Plan")
    add_doc_table(
        doc,
        ["Failure", "What to say", "Recovery"],
        [
            ("Slow model response", "The retrieval is running across multiple signals.", "Use the prepared result screenshot after 15 seconds; continue with citations."),
            ("No proposal appears", "I will use the deterministic simulator path to show the same contract.", "Run the pre-tested CLI prompt or open the saved proposal result."),
            ("Write already exists", "The tool is idempotent, so it will not duplicate an existing action.", "Show the existing row and audit metadata; explain idempotency."),
            ("Persona session is wrong", "Identity is the governance boundary.", "Switch to the pre-opened session; do not edit environment variables live."),
            ("Browser unavailable", "The wire contract has already been validated independently.", "Use screenshots and show capa_tracker.json before/after."),
        ],
        [1.5, 2.8, 3.8],
    )

    add_doc_heading(doc, "Q&A Answers to Keep Ready")
    qa = [
        ("Is this live Work IQ?", "This recording uses the Local Work IQ Simulator. It mirrors the MCP/A2A contract, citation shape, persona trimming, and persistent Tools writes. The production endpoint is a configuration swap after tenant authorization."),
        ("Why both MCP and A2A?", "A2A is the natural fit for delegating a grounded work question; MCP is the fit for explicit structured fetch and governed entity actions."),
        ("Does the agent bypass permissions?", "No. Every production call is delegated as the signed-in user. The local demo models that behavior through persona ACLs."),
        ("Why human approval?", "A CAPA tracker is an operational system of record. The agent may recommend, but a named accountable user approves material writes."),
        ("How do you prevent hallucination?", "The agent must cite source items, withholds unsupported claims, and exposes source IDs for meeting, email, file, and tracker evidence."),
    ]
    for question, answer in qa:
        paragraph = doc.add_paragraph()
        run = paragraph.add_run(question + " ")
        run.bold = True
        paragraph.add_run(answer)

    add_doc_heading(doc, "Final Recording Checklist")
    for item in [
        "The recording is no longer than 15:00.",
        "The Local Work IQ Simulator is disclosed verbally and on the architecture slide.",
        "At least one compound answer shows citations from multiple signal types.",
        "At least one approved Tools write is visible on a follow-up fetch.",
        "The governance demo uses the exact same prompt for both personas.",
        "The close includes quantified impact and a specific pilot decision request.",
    ]:
        add_bullet(doc, "☐ " + item)

    doc.save(DOCX_OUT)


def validate_outputs() -> None:
    deck = Presentation(str(PPTX_OUT))
    if len(deck.slides) != 8:
        raise RuntimeError(f"Expected 8 slides, found {len(deck.slides)}")
    titles = []
    for slide in deck.slides:
        text = " ".join(shape.text for shape in slide.shapes if hasattr(shape, "text_frame"))
        titles.append(text)
    required = ["From committee decisions", "Same question", "Technical deep dive", "90-day pilot"]
    for phrase in required:
        if not any(phrase.casefold() in text.casefold() for text in titles):
            raise RuntimeError(f"Missing required deck content: {phrase}")

    guide = Document(str(DOCX_OUT))
    guide_text = "\n".join(paragraph.text for paragraph in guide.paragraphs)
    for phrase in ["Run of Show", "Approve the PharmaLink", "Demo Recovery Plan", "15:00"]:
        if phrase not in guide_text and not any(
            phrase in cell.text for table in guide.tables for row in table.rows for cell in row.cells
        ):
            raise RuntimeError(f"Missing required guide content: {phrase}")


if __name__ == "__main__":
    build_presentation()
    build_presenter_guide()
    validate_outputs()
    print(f"Created {PPTX_OUT.name} (8 slides)")
    print(f"Created {DOCX_OUT.name}")