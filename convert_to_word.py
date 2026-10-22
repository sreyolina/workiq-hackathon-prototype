"""Convert the Northbridge demo runbook markdown to Word."""
import re
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()
style = doc.styles['Normal']
style.font.name = 'Calibri'
style.font.size = Pt(11)

md_path = Path(__file__).parent / "docs" / "NORTHBRIDGE_DEMO_RUNBOOK.md"
with open(md_path, "r", encoding="utf-8") as f:
    lines = f.readlines()


def flush_table(table_rows, doc):
    if not table_rows:
        return
    headers = [c.strip() for c in table_rows[0].strip("|").split("|")]
    data = []
    for row in table_rows[2:]:
        cells = [c.strip() for c in row.strip("|").split("|")]
        data.append(cells)
    t = doc.add_table(rows=1 + len(data), cols=len(headers))
    t.style = "Table Grid"
    for j, h in enumerate(headers):
        cell_text = re.sub(r"\*\*(.*?)\*\*", r"\1", h)
        t.rows[0].cells[j].text = cell_text
        for p in t.rows[0].cells[j].paragraphs:
            for run in p.runs:
                run.bold = True
    for ri, row_data in enumerate(data):
        for ci, cell in enumerate(row_data):
            if ci < len(t.columns):
                clean = re.sub(r"\*\*(.*?)\*\*", r"\1", cell)
                clean = re.sub(r"`(.*?)`", r"\1", clean)
                t.rows[ri + 1].cells[ci].text = clean
    doc.add_paragraph()


i = 0
in_code_block = False
in_table = False
table_rows_buf = []

while i < len(lines):
    line = lines[i]
    stripped = line.strip()

    # Code block fences
    if stripped.startswith("```"):
        if in_code_block:
            in_code_block = False
        else:
            in_code_block = True
        i += 1
        continue

    if in_code_block:
        p = doc.add_paragraph(line.rstrip())
        p.paragraph_format.left_indent = Inches(0.5)
        for run in p.runs:
            run.font.name = "Consolas"
            run.font.size = Pt(9)
        i += 1
        continue

    # Tables
    if "|" in stripped and stripped.startswith("|"):
        if not in_table:
            in_table = True
            table_rows_buf = []
        table_rows_buf.append(stripped)
        i += 1
        continue
    elif in_table:
        flush_table(table_rows_buf, doc)
        in_table = False
        table_rows_buf = []

    # Horizontal rules
    if stripped in ("---", "***", "___"):
        i += 1
        continue

    # Headings
    if stripped.startswith("# ") and not stripped.startswith("## "):
        doc.add_heading(stripped[2:], level=0)
    elif stripped.startswith("## "):
        doc.add_heading(stripped[3:], level=1)
    elif stripped.startswith("### "):
        doc.add_heading(stripped[4:], level=2)
    elif stripped.startswith("#### "):
        doc.add_heading(stripped[5:], level=3)
    elif stripped.startswith("> "):
        text = re.sub(r"\*\*(.*?)\*\*", r"\1", stripped[2:])
        text = re.sub(r"`(.*?)`", r"\1", text)
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.3)
        run = p.add_run(text)
        run.italic = True
    elif stripped.startswith("- ") or stripped.startswith("* "):
        text = re.sub(r"\*\*(.*?)\*\*", r"\1", stripped[2:])
        text = re.sub(r"`(.*?)`", r"\1", text)
        doc.add_paragraph(text, style="List Bullet")
    elif re.match(r"^\d+\.\s", stripped):
        text = re.sub(r"\*\*(.*?)\*\*", r"\1", stripped)
        text = re.sub(r"`(.*?)`", r"\1", text)
        doc.add_paragraph(text, style="List Number")
    elif stripped == "":
        pass
    else:
        text = re.sub(r"\*\*(.*?)\*\*", r"\1", stripped)
        text = re.sub(r"`(.*?)`", r"\1", text)
        doc.add_paragraph(text)

    i += 1

if in_table:
    flush_table(table_rows_buf, doc)

out_path = Path(__file__).parent / "docs" / "NORTHBRIDGE_DEMO_RUNBOOK.docx"
doc.save(str(out_path))
print(f"Saved: {out_path}")
